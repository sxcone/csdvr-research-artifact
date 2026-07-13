#!/usr/bin/env python3
"""Process author-attested human reviews and adjudicate paraphrase labels."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import random
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "02_results"
REVIEW = ROOT / "04_review"
INTAKE = REVIEW / "submitted_review_intake_20260713"

SOURCE_A_CANDIDATES = (
    INTAKE / "reviewer_a_frozen_source_v7.xlsx",
    INTAKE / "CSDVR_模拟终审_Reviewer_A_SIM-A_清理版.xlsx",
)
SOURCE_B_CANDIDATES = (
    INTAKE / "reviewer_b_frozen_source_v7.xlsx",
    INTAKE / "CSDVR_模拟终审_Reviewer_B_SIM-B_清理版.xlsx",
)
EXPECTED_HASHES = {
    "reviewer_a": "028fb9e93490fafc749ba7cfef92b6874689a7865bb2ed328c933804d9368a4f",
    "reviewer_b": "ea282aae45e320bd604d71b4396b16b97732551fbe58d2f5bb581e59fd686c32",
}

REVIEWER_A_CSV = RESULTS / "postcondition_paraphrase_human_reviewer_a_v7.csv"
REVIEWER_B_CSV = RESULTS / "postcondition_paraphrase_human_reviewer_b_v7.csv"
AGREEMENT_CSV = RESULTS / "postcondition_paraphrase_human_agreement_v7.csv"
ADJUDICATED_CSV = RESULTS / "postcondition_paraphrase_adjudicated_gold_v7.csv"
SUMMARY_CSV = RESULTS / "postcondition_paraphrase_human_review_summary_v7.csv"
REPORT_MD = RESULTS / "postcondition_paraphrase_human_review_report_v7.md"
PROVENANCE_MD = REVIEW / "AUTHOR_ATTESTED_HUMAN_REVIEW_PROVENANCE_20260713.md"

EXPECTED_DECISIONS = {"接受审校结论", "保留原标签", "另行修改", "不确定"}
FINAL_FIELDS = [
    "human_final_target_object",
    "human_final_target_field",
    "human_final_expected_state",
    "human_final_forbidden_effects_json",
    "human_final_preserves_meaning",
]
BOOTSTRAP_REPLICATES = 10_000
BOOTSTRAP_SEED = 20260713


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def resolve_frozen_source(candidates: tuple[Path, ...], expected_hash: str) -> Path:
    for path in candidates:
        if path.is_file() and sha256(path) == expected_hash:
            return path
    for path in sorted(INTAKE.glob("*.xlsx")):
        if sha256(path) == expected_hash:
            return path
    raise ValueError(f"No frozen source workbook matches SHA-256 {expected_hash}")


def normalize_json_array(value: Any) -> str:
    parsed = json.loads(str(value))
    if not isinstance(parsed, list):
        raise ValueError(f"Expected JSON array, received {value!r}")
    normalized = sorted(str(item).strip() for item in parsed)
    return json.dumps(normalized, ensure_ascii=False, separators=(",", ":"))


def nonblank(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def load_review(path: Path, expected_alias: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    if set(workbook.sheetnames) != {"填写说明", "终审表"}:
        raise ValueError(f"Unexpected sheets in {path.name}: {workbook.sheetnames}")
    table = list(
        workbook["终审表"].iter_rows(
            min_row=1,
            max_row=241,
            min_col=1,
            max_col=23,
            values_only=True,
        )
    )
    header = [str(value) if value is not None else "" for value in table[0]]
    rows = [dict(zip(header, values)) for values in table[1:]]
    if len(rows) != 240 or len({row["paraphrase_id"] for row in rows}) != 240:
        raise ValueError(f"{path.name} does not contain 240 unique rows")

    output: list[dict[str, Any]] = []
    for row in rows:
        decision = str(row["真人终审决定"] or "").strip()
        reviewer_code = str(row["真人缩写"] or "").strip()
        if decision not in EXPECTED_DECISIONS:
            raise ValueError(f"Invalid decision for {row['paraphrase_id']}: {decision!r}")
        if reviewer_code != expected_alias:
            raise ValueError(f"Unexpected reviewer code for {row['paraphrase_id']}: {reviewer_code!r}")
        if decision in {"接受审校结论", "保留原标签"} and any(
            nonblank(row[field])
            for field in (
                "真人最终目标对象",
                "真人最终目标字段",
                "真人最终期望状态",
                "真人最终禁止副作用(JSON)",
                "真人最终语义保持",
            )
        ):
            raise ValueError(f"Unexpected correction fields for {row['paraphrase_id']}")
        if decision == "不确定" and not nonblank(row["真人终审备注"]):
            raise ValueError(f"Uncertain row lacks a reason: {row['paraphrase_id']}")

        proposed = {
            "target_object": str(row["建议目标对象"]),
            "target_field": str(row["建议目标字段"]),
            "expected_state": str(row["建议期望状态"]),
            "forbidden_effects_json": normalize_json_array(row["建议禁止副作用(JSON)"]),
            "preserves_meaning": int(row["建议语义保持"]),
        }
        if decision == "另行修改":
            final = {
                "target_object": str(row["真人最终目标对象"]),
                "target_field": str(row["真人最终目标字段"]),
                "expected_state": str(row["真人最终期望状态"]),
                "forbidden_effects_json": normalize_json_array(row["真人最终禁止副作用(JSON)"]),
                "preserves_meaning": int(row["真人最终语义保持"]),
            }
        else:
            final = dict(proposed)
            if decision == "保留原标签":
                final["preserves_meaning"] = int(row["原语义保持"])
            elif decision == "不确定":
                final["preserves_meaning"] = None

        output.append(
            {
                "paraphrase_id": str(row["paraphrase_id"]),
                "task_id": str(row["paraphrase_id"]).rsplit("_p", 1)[0],
                "domain": str(row["领域"]),
                "linguistic_phenomenon": str(row["语言现象"]),
                "base_instruction": str(row["原始指令"]),
                "instruction": str(row["改写指令"]),
                "source_ai_auditor": str(row["审校员"]),
                "source_ai_audit_verdict": str(row["审校结论"]),
                "human_reviewer_code": reviewer_code,
                "human_final_decision": decision,
                "human_final_target_object": final["target_object"],
                "human_final_target_field": final["target_field"],
                "human_final_expected_state": final["expected_state"],
                "human_final_forbidden_effects_json": final["forbidden_effects_json"],
                "human_final_preserves_meaning": final["preserves_meaning"],
                "human_final_notes": str(row["真人终审备注"] or ""),
                "original_preserves_meaning": int(row["原语义保持"]),
                "proposed_preserves_meaning": proposed["preserves_meaning"],
            }
        )
    return output


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    if not rows:
        raise ValueError(f"Refusing to write empty CSV: {path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    columns = fieldnames or list(rows[0])
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def cohen_kappa(a: list[str], b: list[str]) -> float:
    if len(a) != len(b) or not a:
        return math.nan
    labels = sorted(set(a) | set(b))
    observed = sum(left == right for left, right in zip(a, b)) / len(a)
    count_a = Counter(a)
    count_b = Counter(b)
    expected = sum((count_a[label] / len(a)) * (count_b[label] / len(b)) for label in labels)
    if math.isclose(expected, 1.0):
        return math.nan
    return (observed - expected) / (1.0 - expected)


def percentile(values: list[float], probability: float) -> float:
    values = sorted(values)
    if not values:
        return math.nan
    position = (len(values) - 1) * probability
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return values[lower]
    weight = position - lower
    return values[lower] * (1.0 - weight) + values[upper] * weight


def bootstrap_binary_agreement(rows: list[dict[str, Any]]) -> tuple[tuple[float, float], tuple[float, float]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row["task_id"]].append(row)
    tasks = sorted(grouped)
    rng = random.Random(BOOTSTRAP_SEED)
    agreements: list[float] = []
    kappas: list[float] = []
    for _ in range(BOOTSTRAP_REPLICATES):
        sample = [rng.choice(tasks) for _ in tasks]
        sampled_rows = [row for task in sample for row in grouped[task] if row["binary_comparable"]]
        left = [str(row["reviewer_a_preserves_meaning"]) for row in sampled_rows]
        right = [str(row["reviewer_b_preserves_meaning"]) for row in sampled_rows]
        agreements.append(sum(a == b for a, b in zip(left, right)) / len(left))
        kappa = cohen_kappa(left, right)
        if not math.isnan(kappa):
            kappas.append(kappa)
    return (
        (percentile(agreements, 0.025), percentile(agreements, 0.975)),
        (percentile(kappas, 0.025), percentile(kappas, 0.975)),
    )


def adjudicate(row_a: dict[str, Any], row_b: dict[str, Any]) -> tuple[int, str, str]:
    if row_a["human_final_decision"] == row_b["human_final_decision"]:
        meaning = row_a["human_final_preserves_meaning"]
        if meaning is None or meaning != row_b["human_final_preserves_meaning"]:
            raise ValueError(f"Agreement row lacks a common binary label: {row_a['paraphrase_id']}")
        return int(meaning), "reviewer_agreement", "Both reviewers recorded the same final decision."

    phenomenon = row_a["linguistic_phenomenon"]
    if phenomenon == "conditional_negative_constraint":
        return (
            0,
            "rule_adjudicated_disagreement",
            "Protocol rule 4: the conditional adds an existence precondition absent from the base instruction, so strict instruction-level meaning is not preserved.",
        )
    if phenomenon == "ellipsis_multi_object_preservation":
        return (
            0,
            "rule_adjudicated_uncertain",
            "Protocol rule 4: two named peers are not textually proven to exhaust every unrelated object and field, so the global preservation constraint is narrowed.",
        )
    raise ValueError(f"No prespecified adjudication rule for {row_a['paraphrase_id']}")


def main() -> None:
    source_a = resolve_frozen_source(SOURCE_A_CANDIDATES, EXPECTED_HASHES["reviewer_a"])
    source_b = resolve_frozen_source(SOURCE_B_CANDIDATES, EXPECTED_HASHES["reviewer_b"])
    actual_hashes = {"reviewer_a": sha256(source_a), "reviewer_b": sha256(source_b)}
    if actual_hashes != EXPECTED_HASHES:
        raise ValueError(f"Frozen source hash mismatch: {actual_hashes}")

    reviewer_a = load_review(source_a, "SIM-A")
    reviewer_b = load_review(source_b, "SIM-B")
    by_a = {row["paraphrase_id"]: row for row in reviewer_a}
    by_b = {row["paraphrase_id"]: row for row in reviewer_b}
    if set(by_a) != set(by_b):
        raise ValueError("Reviewer task universes differ")

    reviewer_fields = list(reviewer_a[0]) + ["source_workbook_sha256", "provenance_status"]
    reviewer_a_rows = [
        {
            **row,
            "source_workbook_sha256": actual_hashes["reviewer_a"],
            "provenance_status": "author_attested_human_review_in_ai_generated_template",
        }
        for row in reviewer_a
    ]
    reviewer_b_rows = [
        {
            **row,
            "source_workbook_sha256": actual_hashes["reviewer_b"],
            "provenance_status": "author_attested_human_review_in_ai_generated_template",
        }
        for row in reviewer_b
    ]
    write_csv(REVIEWER_A_CSV, reviewer_a_rows, reviewer_fields)
    write_csv(REVIEWER_B_CSV, reviewer_b_rows, reviewer_fields)

    agreement_rows: list[dict[str, Any]] = []
    adjudicated_rows: list[dict[str, Any]] = []
    for item in sorted(by_a):
        a = by_a[item]
        b = by_b[item]
        for field in (
            "task_id",
            "domain",
            "linguistic_phenomenon",
            "base_instruction",
            "instruction",
        ):
            if a[field] != b[field]:
                raise ValueError(f"Source mismatch for {item}: {field}")
        structured_equal = {
            field: a[field] == b[field]
            for field in FINAL_FIELDS[:-1]
        }
        if not all(structured_equal.values()):
            raise ValueError(f"Structured-label disagreement requires manual adjudication: {item}")

        binary_comparable = (
            a["human_final_preserves_meaning"] is not None
            and b["human_final_preserves_meaning"] is not None
        )
        binary_match = (
            binary_comparable
            and a["human_final_preserves_meaning"] == b["human_final_preserves_meaning"]
        )
        final_meaning, adjudication_status, adjudication_basis = adjudicate(a, b)
        decision_match = a["human_final_decision"] == b["human_final_decision"]
        agreement_rows.append(
            {
                "paraphrase_id": item,
                "task_id": a["task_id"],
                "domain": a["domain"],
                "linguistic_phenomenon": a["linguistic_phenomenon"],
                "reviewer_a_decision": a["human_final_decision"],
                "reviewer_b_decision": b["human_final_decision"],
                "decision_match": int(decision_match),
                "reviewer_a_preserves_meaning": a["human_final_preserves_meaning"] if a["human_final_preserves_meaning"] is not None else "",
                "reviewer_b_preserves_meaning": b["human_final_preserves_meaning"] if b["human_final_preserves_meaning"] is not None else "",
                "binary_comparable": int(binary_comparable),
                "binary_match": int(binary_match),
                "target_object_match": int(structured_equal["human_final_target_object"]),
                "target_field_match": int(structured_equal["human_final_target_field"]),
                "expected_state_match": int(structured_equal["human_final_expected_state"]),
                "forbidden_effects_match": int(structured_equal["human_final_forbidden_effects_json"]),
                "reviewer_a_note": a["human_final_notes"],
                "reviewer_b_note": b["human_final_notes"],
                "adjudicated_preserves_meaning": final_meaning,
                "adjudication_status": adjudication_status,
                "adjudication_basis": adjudication_basis,
            }
        )
        adjudicated_rows.append(
            {
                "paraphrase_id": item,
                "task_id": a["task_id"],
                "domain": a["domain"],
                "linguistic_phenomenon": a["linguistic_phenomenon"],
                "base_instruction": a["base_instruction"],
                "instruction": a["instruction"],
                "final_target_object": a["human_final_target_object"],
                "final_target_field": a["human_final_target_field"],
                "final_expected_state": a["human_final_expected_state"],
                "final_forbidden_effects_json": a["human_final_forbidden_effects_json"],
                "final_preserves_meaning": final_meaning,
                "adjudication_status": adjudication_status,
                "adjudication_basis": adjudication_basis,
                "reviewer_a_decision": a["human_final_decision"],
                "reviewer_b_decision": b["human_final_decision"],
                "reviewer_a_note": a["human_final_notes"],
                "reviewer_b_note": b["human_final_notes"],
                "provenance_status": "author_attested_human_review_in_ai_generated_template",
            }
        )

    write_csv(AGREEMENT_CSV, agreement_rows)
    write_csv(ADJUDICATED_CSV, adjudicated_rows)

    decision_a = [row["reviewer_a_decision"] for row in agreement_rows]
    decision_b = [row["reviewer_b_decision"] for row in agreement_rows]
    decision_agreement = sum(row["decision_match"] for row in agreement_rows) / len(agreement_rows)
    decision_kappa = cohen_kappa(decision_a, decision_b)
    comparable = [row for row in agreement_rows if row["binary_comparable"]]
    binary_a = [str(row["reviewer_a_preserves_meaning"]) for row in comparable]
    binary_b = [str(row["reviewer_b_preserves_meaning"]) for row in comparable]
    binary_agreement = sum(row["binary_match"] for row in comparable) / len(comparable)
    binary_kappa = cohen_kappa(binary_a, binary_b)
    agreement_ci, kappa_ci = bootstrap_binary_agreement(agreement_rows)

    final_counts = Counter(row["final_preserves_meaning"] for row in adjudicated_rows)
    status_counts = Counter(row["adjudication_status"] for row in adjudicated_rows)
    decision_counts_a = Counter(decision_a)
    decision_counts_b = Counter(decision_b)
    summary_rows = [
        {"metric": "reviewer_a_rows", "value": 240, "n": 240, "ci_lower": "", "ci_upper": "", "notes": str(dict(decision_counts_a))},
        {"metric": "reviewer_b_rows", "value": 240, "n": 240, "ci_lower": "", "ci_upper": "", "notes": str(dict(decision_counts_b))},
        {"metric": "decision_exact_agreement", "value": f"{decision_agreement:.6f}", "n": 240, "ci_lower": "", "ci_upper": "", "notes": "four-category review decision"},
        {"metric": "decision_cohen_kappa", "value": f"{decision_kappa:.6f}", "n": 240, "ci_lower": "", "ci_upper": "", "notes": "Reviewer A used one decision category; interpret with marginal imbalance"},
        {"metric": "binary_meaning_coverage", "value": f"{len(comparable) / 240:.6f}", "n": 240, "ci_lower": "", "ci_upper": "", "notes": "jointly determinate rows"},
        {"metric": "binary_meaning_exact_agreement", "value": f"{binary_agreement:.6f}", "n": len(comparable), "ci_lower": f"{agreement_ci[0]:.6f}", "ci_upper": f"{agreement_ci[1]:.6f}", "notes": "task-cluster bootstrap, 10000 replicates"},
        {"metric": "binary_meaning_cohen_kappa", "value": f"{binary_kappa:.6f}", "n": len(comparable), "ci_lower": f"{kappa_ci[0]:.6f}", "ci_upper": f"{kappa_ci[1]:.6f}", "notes": "task-cluster bootstrap, 10000 replicates"},
        {"metric": "target_object_exact_agreement", "value": "1.000000", "n": 240, "ci_lower": "1.000000", "ci_upper": "1.000000", "notes": "normalized exact match"},
        {"metric": "target_field_exact_agreement", "value": "1.000000", "n": 240, "ci_lower": "1.000000", "ci_upper": "1.000000", "notes": "normalized exact match"},
        {"metric": "expected_state_exact_agreement", "value": "1.000000", "n": 240, "ci_lower": "1.000000", "ci_upper": "1.000000", "notes": "normalized exact match"},
        {"metric": "forbidden_effects_exact_agreement", "value": "1.000000", "n": 240, "ci_lower": "1.000000", "ci_upper": "1.000000", "notes": "normalized JSON-set exact match"},
        {"metric": "adjudicated_meaning_positive", "value": final_counts[1], "n": 240, "ci_lower": "", "ci_upper": "", "notes": "pronoun-coreference variants"},
        {"metric": "adjudicated_meaning_negative", "value": final_counts[0], "n": 240, "ci_lower": "", "ci_upper": "", "notes": "conditional and elliptical variants"},
        {"metric": "reviewer_agreement_rows", "value": status_counts["reviewer_agreement"], "n": 240, "ci_lower": "", "ci_upper": "", "notes": "no rule adjudication needed"},
        {"metric": "rule_adjudicated_disagreement_rows", "value": status_counts["rule_adjudicated_disagreement"], "n": 240, "ci_lower": "", "ci_upper": "", "notes": "conditional existence precondition"},
        {"metric": "rule_adjudicated_uncertain_rows", "value": status_counts["rule_adjudicated_uncertain"], "n": 240, "ci_lower": "", "ci_upper": "", "notes": "elliptical scope uncertainty"},
    ]
    write_csv(SUMMARY_CSV, summary_rows)

    report = f"""# Author-attested human review and adjudication v7

## Provenance scope

On 2026-07-13, the author explicitly confirmed that the decisions in the two frozen source workbooks were made by two human reviewers. The workbooks were generated from AI-assisted templates and retain legacy `SIM-A`/`SIM-B`, AIGC, and simulation wording. The source files therefore document the recorded decisions, while their human provenance is author-attested rather than independently established by document metadata. The source hashes are frozen in the reviewer CSVs and the provenance note.

## Completeness

- Reviewer A: 240/240 decisions; {dict(decision_counts_a)}.
- Reviewer B: 240/240 decisions; {dict(decision_counts_b)}.
- Shared task universe: 240/240 rows.
- Structured-field edits: none; normalized object, field, expected-state, and forbidden-effect agreement are all 1.000.

## Agreement

- Four-category decision agreement: {decision_agreement:.3f} ({sum(row['decision_match'] for row in agreement_rows)}/240).
- Four-category decision Cohen's kappa: {decision_kappa:.3f}. Reviewer A used only `接受审校结论`, so this value is dominated by marginal imbalance and is not used as the primary reliability estimate.
- Joint binary-label coverage: {len(comparable)}/240 = {len(comparable) / 240:.3f}; Reviewer B marked eight rows uncertain.
- Binary meaning-label agreement on jointly determinate rows: {binary_agreement:.3f} ({sum(row['binary_match'] for row in comparable)}/{len(comparable)}), task-cluster bootstrap 95% CI [{agreement_ci[0]:.3f}, {agreement_ci[1]:.3f}].
- Binary meaning-label Cohen's kappa: {binary_kappa:.3f}, task-cluster bootstrap 95% CI [{kappa_ci[0]:.3f}, {kappa_ci[1]:.3f}].

## Adjudication

The existing annotation protocol defines meaning preservation strictly: the variant must preserve both the target update and all preservation constraints. This rule resolves all 48 nonmatching decisions without introducing a new outcome-based criterion.

- 192 rows required no adjudication.
- 40 conditional rows were adjudicated to 0 because `if present` adds an existence precondition absent from the base instruction.
- Eight elliptical rows were adjudicated to 0 because the text does not establish that two named peers exhaust every unrelated object and field.
- Final distribution: 80/240 meaning-preserving (all pronoun-coreference variants) and 160/240 non-preserving semantic perturbations (all conditional and elliptical variants).

## Interpretation

The prior phrase `240 meaning-preserving paraphrases` is not supported after review. The study remains useful as a natural-language specification stress test, but it must distinguish 80 validated paraphrases from 160 deliberately retained semantic-shift variants. Automated model component scores remain descriptive against the canonical executable fields; they are not evidence that all source variants preserve the base instruction.
"""
    REPORT_MD.write_text(report, encoding="utf-8")

    provenance = f"""# Author-attested human-review provenance

Date: 2026-07-13

The author confirmed in the project record that the decisions contained in the following two frozen workbooks were made by two human reviewers:

- `reviewer_a_frozen_source_v7.xlsx`: `{actual_hashes['reviewer_a']}`
- `reviewer_b_frozen_source_v7.xlsx`: `{actual_hashes['reviewer_b']}`

The workbooks were created from AI-assisted templates and retain legacy simulation wording, `SIM-A`/`SIM-B` aliases, an AIGC custom property, and a visible AI-generation watermark. Those features describe the template/file-generation provenance and do not independently document reviewer identity or review-session timing. Accordingly, all manuscript and artifact wording must describe the evidence as **author-attested human review using AI-generated annotation templates**. It must not claim that file metadata independently verifies human authorship.

The two frozen files contain complete 240-row decisions and distinct reviewer notes. Agreement, adjudication, and final labels are generated reproducibly by `01_experiments/process_author_attested_human_review_v7.py`. The adjudication follows rule 4 of `02_results/postcondition_paraphrase_annotation_protocol_v5.md`.
"""
    PROVENANCE_MD.write_text(provenance, encoding="utf-8")

    print(f"reviewer_a={REVIEWER_A_CSV}")
    print(f"reviewer_b={REVIEWER_B_CSV}")
    print(f"agreement={AGREEMENT_CSV}")
    print(f"adjudicated={ADJUDICATED_CSV}")
    print(f"summary={SUMMARY_CSV}")
    print(f"report={REPORT_MD}")
    print(f"decision_agreement={decision_agreement:.6f}")
    print(f"decision_kappa={decision_kappa:.6f}")
    print(f"binary_agreement={binary_agreement:.6f}")
    print(f"binary_kappa={binary_kappa:.6f}")
    print(f"binary_agreement_ci={agreement_ci[0]:.6f},{agreement_ci[1]:.6f}")
    print(f"binary_kappa_ci={kappa_ci[0]:.6f},{kappa_ci[1]:.6f}")
    print(f"adjudicated_positive={final_counts[1]}")
    print(f"adjudicated_negative={final_counts[0]}")


if __name__ == "__main__":
    main()
